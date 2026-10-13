"""Perguntas ao cliente: a planilha com o que o sistema não reconheceu, e as respostas voltando para a Revisão,
para o histórico do lançamento e para a IA."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.test_classificacao_ia import responder
from tests.test_revisao import empresa, recarregar, sugestao_de  # noqa: F401 — empresa com o extrato de agosto

pytestmark = pytest.mark.banco


def baixar(cliente, empresa):  # noqa: F811
    cliente.get(f"/empresas/{empresa.id}", params={"aba": "revisao", "mes": "2026-08"})  # o mês fica na tela
    resposta = cliente.get(f"/empresas/{empresa.id}/perguntas.xlsx")
    assert 'filename="perguntas_230_2026-08.xlsx"' in resposta.headers["content-disposition"]
    return load_workbook(BytesIO(resposta.content))


def linha_de(folha, texto):
    return next(linha for linha in folha.iter_rows(min_row=5) if texto in str(linha[3].value))


def test_planilha_traz_so_o_que_o_sistema_nao_reconheceu(cliente, empresa, sessao):  # noqa: F811
    folha = baixar(cliente, empresa).active
    assert folha["A4"].value == "Ref." and folha["G4"].value == "O que é? (responda aqui)"
    descricoes = [str(linha[3].value) for linha in folha.iter_rows(min_row=5)]
    assert any("FULANO DESCONHECIDO" in d for d in descricoes)
    assert not any("COMPANHIA DE AGUA" in d for d in descricoes)  # reconhecida (verde): não pergunta
    fulano = linha_de(folha, "FULANO DESCONHECIDO")
    assert fulano[0].value == sugestao_de(sessao, "FULANO DESCONHECIDO").id and fulano[4].value == -77


def responder_fulano(cliente, empresa, resposta="Material de limpeza; mercado"):  # noqa: F811
    livro = baixar(cliente, empresa)
    linha_de(livro.active, "FULANO DESCONHECIDO")[6].value = resposta
    arquivo = BytesIO()
    livro.save(arquivo)
    return cliente.post(f"/empresas/{empresa.id}/perguntas", files={"arquivo": ("r.xlsx", arquivo.getvalue())}).text


def test_resposta_vai_para_a_revisao_e_o_historico(cliente, empresa, sessao):  # noqa: F811
    tela = responder_fulano(cliente, empresa)
    assert "Respostas do cliente: 1 linha(s) receberam a resposta." in tela
    assert "💬 Cliente: Material de limpeza, mercado" in tela  # o “;” virou “,” (o Domínio usa “;”)
    fulano = recarregar(sessao, sugestao_de(sessao, "FULANO DESCONHECIDO"))
    assert fulano.resposta_cliente == "Material de limpeza, mercado" and fulano.status == "PENDENTE"
    assert "MATERIAL DE LIMPEZA, MERCADO" in fulano.complemento  # no lugar da natureza do histórico


def test_ia_recebe_a_resposta_do_cliente(cliente, empresa, sessao, ia_falsa):  # noqa: F811
    ia_falsa.respostas = responder({'o cliente disse que é: "Material de limpeza, mercado"': (359, 92)})
    responder_fulano(cliente, empresa)
    fulano = recarregar(sessao, sugestao_de(sessao, "FULANO DESCONHECIDO"))
    assert (fulano.conta_debito.reduzido, fulano.origem) == (359, "IA")  # a IA usou a resposta


def test_planilha_errada_ou_linha_ja_decidida(cliente, empresa, sessao):  # noqa: F811
    tela = cliente.post(f"/empresas/{empresa.id}/perguntas", files={"arquivo": ("x.xlsx", b"nao e planilha")}).text
    assert "Não consegui abrir a planilha" in tela
    fulano = sugestao_de(sessao, "FULANO DESCONHECIDO")
    livro = baixar(cliente, empresa)
    cliente.post(f"/empresas/{empresa.id}/revisao/{fulano.id}", data={"acao": "salvar", "conta": "359"})
    linha_de(livro.active, "FULANO DESCONHECIDO")[6].value = "chegou tarde"
    arquivo = BytesIO()
    livro.save(arquivo)
    tela = cliente.post(f"/empresas/{empresa.id}/perguntas", files={"arquivo": ("r.xlsx", arquivo.getvalue())}).text
    assert "1 já tinham sido decididas" in tela and "Nenhuma resposta nova" in tela
    assert recarregar(sessao, fulano).resposta_cliente is None
