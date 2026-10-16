"""Demonstrativo de prestação de contas de teste, no leiaute do real (aba DEMOSTRATIVO, blocos 1 a 5), com dados
fictícios. Colunas como nas creches: A item, B grupo, D favorecido, K especificação, N tipo, O número, R data do
documento, S nº da transferência, T data do pagamento, W valor — e, como nas planilhas reais, uma lista de apoio na
coluna Y e o cabeçalho repetido no meio do bloco 4."""

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook

COLUNAS = {"item": 1, "grupo": 2, "favorecido": 4, "especificacao": 11, "tipo": 14, "numero": 15, "data_doc": 18,
           "transf": 19, "data": 20, "valor": 23}  # fmt: skip


def _cabecalho(aba, linha: int) -> int:
    aba.cell(linha, 1, "11 - ÍTEM")
    aba.cell(linha, 4, "12 - NOME DO FAVORECIDO/CNPJ OU CPF")
    aba.cell(linha, 11, "13 - ESPECIFICAÇÃO DO ")
    aba.cell(linha, 14, "14 - DOCUMENTOS")
    aba.cell(linha, 15, "NÚMERO")
    aba.cell(linha, 18, "DATA")
    aba.cell(linha, 19, "15 - PAGAMENTO")
    aba.cell(linha, 23, "16 - VALOR")
    aba.cell(linha + 1, 11, "PAGAMENTO / CARGO ")
    aba.cell(linha + 2, 11, "MATERIAL OU SERVIÇOS")
    aba.cell(linha + 2, 14, "TIPO")
    aba.cell(linha + 2, 19, "TRANSF. Nº")
    aba.cell(linha + 2, 20, "DATA")
    return linha + 3


def montar_demonstrativo(
    pagamentos: list[dict],
    receitas: list[tuple[str, float]] | None = None,
    conta: str = "AG: 1234-5 C.C Nº: 12345-6",
    periodo: str = "08º MÊS/2026",
    repasse: tuple[datetime, int] = (datetime(2026, 8, 11), 556882000051019),
) -> bytes:
    """pagamentos: [{"favorecido", "especificacao", "tipo", "numero", "data_doc", "transf", "data", "valor",
    "grupo"}]; receitas: [(rótulo do bloco 3, valor)] (“REPASSE MUNICIPAL”, “OUTRAS RECEITAS - … - DOC …”)."""
    livro = Workbook()
    livro.active.title = "DADOS"
    livro["DADOS"]["C3"] = "NOME:"
    aba = livro.create_sheet("DEMOSTRATIVO")
    aba["A1"] = "BLOCO 1 - IDENTIFICAÇÃO / ENTIDADE"
    aba["A2"], aba["D2"] = "1 - NOME DA ENTIDADE:", "Associação Exemplo"
    aba["A3"], aba["B3"] = "2 - CNPJ: ", "11.222.333/0001-81"
    aba["A7"] = "BLOCO 2 - IDENTIFICAÇÃO / CRECHE"
    aba["A8"], aba["D8"], aba["O8"], aba["S8"] = (
        "1 - NOME DA CRECHE:",
        "Creche Exemplo",
        "03 - AG / CONTA",
        "04 - PERIODO",
    )
    aba["O9"], aba["S9"] = conta, periodo
    aba["A13"] = "BLOCO 3 - RECEITAS"
    aba["A14"], aba["D14"], aba["K14"], aba["N14"], aba["T14"] = (
        "DATA DE REPASSE:",
        repasse[0],
        "NUMERO DE TRANSFERENCIA:",
        repasse[1],
        " VALOR ( R$ )",
    )
    aba["A15"], aba["T15"] = "SALDO ANTERIOR: ", 1000.0
    linha = 16
    for rotulo, valor in receitas or [("REPASSE MUNICIPAL: ", 90000.0), ("RENDIMENTO APLICAÇÃO FINANCEIRA:", 10.5)]:
        aba.cell(linha, 1, rotulo)
        aba.cell(linha, 20, valor)
        linha += 1
    aba.cell(linha, 1, "TOTAL")
    aba.cell(linha, 29, "REF S FOLHA DE ABRIL")  # anotação solta à direita, como nas planilhas reais
    linha += 1
    aba.cell(linha, 1, "BLOCO 4 - PAGAMENTOS EFETUADOS")
    linha = _cabecalho(aba, linha + 1)
    total = 0.0
    for numero, pagamento in enumerate(pagamentos, start=1):
        if pagamento.get("cabecalho_antes"):  # o cabeçalho repete na página seguinte
            linha = _cabecalho(aba, linha + 1)
        aba.cell(linha, 1, numero)
        if pagamento.get("grupo"):
            aba.cell(linha, 2, pagamento["grupo"])
        aba.cell(linha, 4, pagamento["favorecido"])
        aba.cell(linha, 11, pagamento["especificacao"])
        aba.cell(linha, 14, pagamento.get("tipo", "NOTA FISCAL"))
        aba.cell(linha, 15, pagamento.get("numero", "S/N"))
        aba.cell(linha, 18, pagamento.get("data_doc"))
        aba.cell(linha, 19, pagamento.get("transf"))
        aba.cell(linha, 20, pagamento["data"])
        aba.cell(linha, 23, pagamento["valor"])
        aba.cell(linha, 25, "LISTA DE APOIO")  # coluna Y: fica de fora
        total += pagamento["valor"]
        linha += 1
    aba.cell(linha, 25, "TOTAL")
    aba.cell(linha, 26, 999999.99)  # o “TOTAL” da conciliação, à direita: não é o total das despesas
    linha += 1
    aba.cell(linha, 1, "TOTAL DAS DESPESAS ")
    aba.cell(linha, 19, "25- TOTAL")
    aba.cell(linha, 23, round(total, 2))
    aba.cell(linha + 1, 1, "BLOCO 5 - SINTESE DA RECEITA E DA DESPESAS")
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
