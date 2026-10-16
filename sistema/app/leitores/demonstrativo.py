"""Demonstrativo da execução de receitas, despesas e pagamentos efetuados — o Excel da prestação de contas das
creches subvencionadas que o escritório monta todo mês (pedido de 15/10/2026).

Aba “DEMOSTRATIVO” (ou outra que tenha o BLOCO 4):
- BLOCO 2: nome da creche, “AG: 1819-8 C.C Nº: 83431-9” (às vezes duas contas: “53215-0 e 73000-9”) e o período
  (“08º MÊS/2026”);
- BLOCO 3 – RECEITAS: data e nº da transferência do repasse, saldo anterior, repasse, rendimento e outras receitas
  (“OUTRAS RECEITAS - TRANSF. RECEBIDA DA CRECHE … - DOC 611.819.000.078.640 - 14/08/2026”);
- BLOCO 4 – PAGAMENTOS EFETUADOS: item, grupo (Recursos Humanos, Encargos…), favorecido/CNPJ, especificação,
  documento (tipo, número, data), pagamento (nº da transferência, data) e valor;
- “TOTAL DAS DESPESAS”, que confere a soma dos pagamentos.
As colunas são achadas pelos títulos (cada creche mexe um pouco no modelo: na Jandira o nº do documento às vezes cai
uma coluna à direita); o que fica à direita do VALOR (listas de apoio que as creches deixam ali) é ignorado.
"""

import re
import warnings
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from typing import Any

from app.servicos.contas_chave import normalizar
from app.servicos.documentos import normalizar_documento, tipo_documento
from app.servicos.plano import Aviso, ErroLeitura

_CNPJ = re.compile(r"\b\d{2}\.?\d{3}\.?\d{3}\s*/?\s*\d{3,4}\s*-?\s*\d{2}\b|\b\d{3}\.\d{3}\.\d{3}-\d{2}\b")
_CONTA = re.compile(r"\d[\d.]*\d-[\dXx]")  # “83431-9”, “72.637-0”
_MES = re.compile(r"(\d{1,2})\s*[º°ªO]?\s*M[EÊ]S\s*/\s*(\d{4})", re.IGNORECASE)
_DATA_NO_TEXTO = re.compile(r"\b(\d{2})/(\d{2})/(\d{4})\b")
_SEM_NUMERO = {"S/N", "SN", "S N", "-", ""}
MAXIMO_LINHAS = 400
MAXIMO_COLUNAS = 60


@dataclass
class LinhaLida:
    """Um pagamento (SAÍDA) ou uma receita (ENTRADA) do demonstrativo."""

    ordem: int
    sentido: str  # SAIDA (pagamento) ou ENTRADA (receita)
    valor: Decimal
    favorecido: str = ""  # como está escrito (“SIEMACO 03.491.527/0001-54”): vai assim para o histórico
    nome: str = ""  # o favorecido sem o CNPJ
    cnpj_cpf: str | None = None
    especificacao: str = ""  # o que foi pago (“SALÁRIO MENSAL”, “MAT. DE LIMPEZA (SABÃO)”)
    grupo: str = ""  # Recursos Humanos, Encargos, Consumo, Materiais, Serviços
    item: int | None = None
    documento_tipo: str = ""  # NOTA FISCAL, HOLERITE, BOLETO, DARF…
    documento_numero: str = ""  # vazio = S/N
    documento_data: date | None = None
    transferencia: str = ""  # nº da transferência/documento do banco, só os dígitos
    data: date | None = None  # data do pagamento (ou do repasse)


@dataclass
class DemonstrativoLido:
    creche: str = ""
    cnpj: str | None = None
    agencia: str = ""
    contas: list[str] = field(default_factory=list)  # “83431-9” (às vezes duas)
    mes: date | None = None  # 1º dia do mês do demonstrativo
    saldo_anterior: Decimal | None = None
    repasse: Decimal | None = None
    rendimento: Decimal | None = None
    total_despesas: Decimal | None = None
    linhas: list[LinhaLida] = field(default_factory=list)
    avisos: list[Aviso] = field(default_factory=list)

    @property
    def pagamentos(self) -> list[LinhaLida]:
        return [linha for linha in self.linhas if linha.sentido == "SAIDA"]

    @property
    def receitas(self) -> list[LinhaLida]:
        return [linha for linha in self.linhas if linha.sentido == "ENTRADA"]


# ---------------------------------------------------------------- células


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return " ".join(str(valor).split())


def _dinheiro(valor: Any) -> Decimal | None:
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, int | float | Decimal):
        return Decimal(str(valor)).quantize(Decimal("0.01"))
    texto = str(valor).replace("R$", "").strip()
    if not re.fullmatch(r"-?[\d.]+,\d{1,2}|-?\d+(\.\d{1,2})?", texto):
        return None
    try:
        return Decimal(texto.replace(".", "").replace(",", ".") if "," in texto else texto).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _data(valor: Any) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if m := _DATA_NO_TEXTO.search(_texto(valor)):
        try:
            return date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def _digitos(valor: Any) -> str:
    return re.sub(r"\D", "", _texto(valor))


class _Planilha:
    """As células com conteúdo, por linha, de uma aba."""

    def __init__(self, aba):
        self.nome = aba.title
        self.celulas: dict[int, dict[int, Any]] = {}
        for linha in aba.iter_rows(min_row=1, max_row=MAXIMO_LINHAS, max_col=MAXIMO_COLUNAS):
            for celula in linha:
                if celula.value is not None and _texto(celula.value):
                    self.celulas.setdefault(celula.row, {})[celula.column] = celula.value

    def linhas(self) -> list[int]:
        return sorted(self.celulas)

    def textos(self, linha: int) -> dict[int, str]:
        return {coluna: normalizar(_texto(valor)) for coluna, valor in self.celulas.get(linha, {}).items()}

    def linha_com(self, *palavras: str, depois: int = 0) -> int | None:
        for linha in self.linhas():
            if linha > depois and all(any(p in t for t in self.textos(linha).values()) for p in palavras):
                return linha
        return None

    def a_direita(self, linha: int, coluna: int) -> list[tuple[int, Any]]:
        return sorted((c, v) for c, v in self.celulas.get(linha, {}).items() if c > coluna)


# ---------------------------------------------------------------- leitura


def e_demonstrativo(conteudo: bytes) -> bool:
    """O Excel é um demonstrativo de prestação de contas (e não um extrato em planilha)?"""
    if conteudo[:2] != b"PK":
        return False
    try:
        return _aba_do_demonstrativo(conteudo) is not None
    except Exception:  # planilha que nem abre: quem decide é o leitor de planilhas
        return False


def _aba_do_demonstrativo(conteudo: bytes):
    from openpyxl import load_workbook

    with warnings.catch_warnings():  # “Data Validation extension is not supported”: as listas de escolha das creches
        warnings.simplefilter("ignore", UserWarning)
        livro = load_workbook(BytesIO(conteudo), data_only=True)
    for aba in livro.worksheets:
        for linha in aba.iter_rows(min_row=1, max_row=120, max_col=30, values_only=True):
            texto = normalizar(" ".join(_texto(v) for v in linha if v is not None))
            if "PAGAMENTOS EFETUADOS" in texto:
                return aba
    return None


def ler_demonstrativo(nome: str, conteudo: bytes) -> DemonstrativoLido:
    try:
        aba = _aba_do_demonstrativo(conteudo)
    except Exception as erro:
        raise ErroLeitura(f"Não consegui abrir “{nome}” ({erro}). Salve de novo no Excel (.xlsx) e envie.") from erro
    if aba is None:
        raise ErroLeitura(
            f"“{nome}” não tem o “BLOCO 4 - PAGAMENTOS EFETUADOS” do demonstrativo. Confira se é o arquivo certo."
        )
    planilha = _Planilha(aba)
    lido = DemonstrativoLido()
    _identificacao(planilha, lido)
    _receitas(planilha, lido)
    _pagamentos(planilha, lido)
    if lido.mes is None:
        datas = [linha.data for linha in lido.linhas if linha.data]
        if datas:
            meio = sorted(datas)[len(datas) // 2]
            lido.mes = meio.replace(day=1)
            lido.avisos.append(
                Aviso(f"Não achei o “PERÍODO” no demonstrativo: usei o mês dos pagamentos ({meio:%m/%Y}).")
            )
    if not lido.pagamentos:
        raise ErroLeitura(f"Não encontrei pagamentos no BLOCO 4 de “{nome}”. Confira se a planilha está preenchida.")
    if lido.mes is None:
        raise ErroLeitura(f"Não achei o período (ex.: “08º MÊS/2026”) em “{nome}”.")
    return lido


def _identificacao(planilha: _Planilha, lido: DemonstrativoLido) -> None:
    for linha in planilha.linhas():
        for coluna, texto in planilha.textos(linha).items():
            original = _texto(planilha.celulas[linha][coluna])
            if not lido.contas and "AG" in texto and ("C.C" in texto or "CONTA" in texto):
                if m := re.search(r"AG\.?\s*:?\s*(\d{3,5}-?[\dXx]?)", original, re.IGNORECASE):
                    resto = original[m.end() :]
                    contas = [c.replace(".", "") for c in _CONTA.findall(resto)]
                    if contas:
                        lido.agencia, lido.contas = m.group(1), contas
            if lido.mes is None and (m := _MES.search(original)):
                mes, ano = int(m.group(1)), int(m.group(2))
                if 1 <= mes <= 12:
                    lido.mes = date(ano, mes, 1)
            if not lido.creche and "NOME DA CRECHE" in texto:
                lido.creche = next((_texto(v) for _, v in planilha.a_direita(linha, coluna) if _texto(v)), "")
            if lido.cnpj is None and re.search(r"CNPJ:?$", texto):
                for _, valor in planilha.a_direita(linha, coluna):
                    documento = normalizar_documento(_texto(valor))
                    if documento and tipo_documento(documento):
                        lido.cnpj = documento
                        break


def _receitas(planilha: _Planilha, lido: DemonstrativoLido) -> None:
    inicio = planilha.linha_com("BLOCO 3")
    fim = planilha.linha_com("BLOCO 4")
    if inicio is None or fim is None:
        return
    data_repasse, numero_repasse = None, ""
    linhas = [n for n in planilha.linhas() if inicio < n < fim]
    # só até a coluna do VALOR (à direita as creches deixam anotações e listas de apoio)
    limite = next((c + 2 for n in linhas for c, t in sorted(planilha.textos(n).items()) if "VALOR" in t), 30)
    for linha in linhas:
        textos = {c: t for c, t in planilha.textos(linha).items() if c <= limite}
        if not textos:
            continue
        rotulo = textos[min(textos)]
        valores = [(c, v) for c, v in planilha.celulas[linha].items() if min(textos) < c <= limite]
        dinheiro = [d for _, v in sorted(valores) if (d := _dinheiro(v)) is not None and not _e_numero_documento(v)]
        valor = dinheiro[-1] if dinheiro else None
        if "DATA DE REPASSE" in rotulo:
            data_repasse = next((d for _, v in sorted(valores) if (d := _data(v))), None)
            numero_repasse = next(
                (_digitos(v) for _, v in sorted(valores) if not _data(v) and len(_digitos(v)) >= 9), ""
            )
        elif "SALDO ANTERIOR" in rotulo:
            lido.saldo_anterior = valor
        elif "REPASSE" in rotulo:
            lido.repasse = valor
            if valor:
                lido.linhas.append(
                    LinhaLida(
                        len(lido.linhas),
                        "ENTRADA",
                        valor,
                        especificacao="REPASSE MUNICIPAL",
                        transferencia=numero_repasse,
                        data=data_repasse,
                    )  # fmt: skip
                )
        elif "RENDIMENTO" in rotulo:
            lido.rendimento = valor
        elif "OUTRAS RECEITAS" in rotulo and valor:
            original = _texto(planilha.celulas[linha][min(textos)])
            descricao = re.sub(r"^\s*OUTRAS RECEITAS\s*[-:]?\s*", "", original, flags=re.IGNORECASE)
            documento = re.search(r"DOC\.?\s*-?\s*(\d[\d.]{5,}\d)", descricao, re.IGNORECASE)
            lido.linhas.append(
                LinhaLida(
                    len(lido.linhas),
                    "ENTRADA",
                    valor,
                    especificacao=descricao[:200],
                    transferencia=_digitos(documento.group(1)) if documento else "",
                    data=_data(descricao),
                )
            )


def _e_numero_documento(valor: Any) -> bool:
    """Número de transferência (556882000051019) não é dinheiro."""
    return isinstance(valor, int | float) and float(valor).is_integer() and abs(valor) >= 10**8


@dataclass
class _Colunas:
    item: int
    favorecido: int
    especificacao: int
    tipo: int
    numero: int
    data_documento: int
    transferencia: int
    data_pagamento: int
    valor: int

    def faixa(self, campo: str) -> range:
        """Colunas do campo: da dele até antes da próxima (o VALOR só tem a dele e a seguinte)."""
        inicio = getattr(self, campo)
        if campo == "valor":
            return range(inicio, inicio + 2)
        seguintes = [getattr(self, c) for c in self.__dataclass_fields__ if getattr(self, c) > inicio]
        return range(inicio, min(seguintes) if seguintes else inicio + 1)


def _colunas(planilha: _Planilha, linha: int) -> _Colunas | None:
    rotulos: list[tuple[int, str]] = []
    for n in (linha, linha + 1, linha + 2):
        rotulos += sorted(planilha.textos(n).items())

    def coluna(*palavras: str, depois: int = 0) -> int | None:
        return next((c for c, t in sorted(rotulos) if c > depois and any(p in t for p in palavras)), None)

    favorecido = coluna("FAVORECIDO")
    especificacao = coluna("ESPECIFICA")
    tipo = coluna("TIPO")
    numero = coluna("NUMERO", "NUMERO DO DOC")
    transferencia = coluna("TRANSF")
    valor = coluna("VALOR", depois=transferencia or 0)
    if None in (favorecido, especificacao, tipo, transferencia, valor):
        return None
    datas = sorted({c for c, t in rotulos if t == "DATA" or t.endswith(" DATA")})
    data_documento = next((c for c in datas if c < transferencia), (numero or tipo) + 1)
    data_pagamento = next((c for c in datas if c > transferencia), transferencia + 1)
    return _Colunas(
        item=coluna("ITEM") or 1,
        favorecido=favorecido,
        especificacao=especificacao,
        tipo=tipo,
        numero=numero or tipo + 1,
        data_documento=data_documento,
        transferencia=transferencia,
        data_pagamento=data_pagamento,
        valor=valor,
    )


def _pagamentos(planilha: _Planilha, lido: DemonstrativoLido) -> None:
    bloco = planilha.linha_com("PAGAMENTOS EFETUADOS")
    cabecalho = planilha.linha_com("FAVORECIDO", "VALOR", depois=bloco or 0)
    if cabecalho is None or (colunas := _colunas(planilha, cabecalho)) is None:
        raise ErroLeitura(
            "Não achei os títulos do BLOCO 4 (favorecido, especificação, documentos, pagamento, valor) no "
            "demonstrativo. Confira se o modelo é o da prestação de contas."
        )
    grupo = ""
    soma = Decimal("0")
    for linha in [n for n in planilha.linhas() if n > cabecalho]:
        limite = colunas.valor + 2  # à direita do VALOR ficam listas de apoio das creches (“TOTAL” da conciliação)
        celulas = {c: v for c, v in planilha.celulas[linha].items() if c < limite}
        textos = {c: t for c, t in planilha.textos(linha).items() if c < limite}
        if not celulas:
            continue
        primeiro = textos[min(textos)] if textos else ""
        if primeiro.startswith("TOTAL"):
            lido.total_despesas = next(
                (d for c, v in sorted(celulas.items()) if c in colunas.faixa("valor") and (d := _dinheiro(v))), None
            )
            break
        if "BLOCO 5" in primeiro:
            break
        if any("FAVORECIDO" in t or t in ("TIPO", "MATERIAL OU SERVICOS") for t in textos.values()):
            continue  # o cabeçalho repete na página seguinte

        def campo(nome: str, celulas: dict[int, Any] = celulas) -> list[Any]:
            return [celulas[c] for c in colunas.faixa(nome) if c in celulas]

        valor = next((d for v in campo("valor") if (d := _dinheiro(v)) is not None), None)
        if valor is None or valor <= 0:
            continue
        for c in range(colunas.item + 1, colunas.favorecido):  # o grupo fica entre o item e o favorecido
            if c in celulas and _texto(celulas[c]):
                grupo = _texto(celulas[c])
        favorecido = " ".join(_texto(v) for v in campo("favorecido"))
        item = next((int(v) for v in campo("item") if isinstance(v, int | float)), None)
        numero = next((_texto(v) for v in campo("numero") if _texto(v)), "")
        cnpjs = [normalizar_documento(m.group(0)) for m in _CNPJ.finditer(favorecido)]
        cnpj = next((d for d in cnpjs if tipo_documento(d)), None)
        nome = " ".join(_CNPJ.sub(" ", favorecido).split()).strip(" -")
        lido.linhas.append(
            LinhaLida(
                ordem=len(lido.linhas),
                sentido="SAIDA",
                valor=valor,
                favorecido=favorecido[:200],
                nome=nome[:200],
                cnpj_cpf=cnpj,
                especificacao=" ".join(_texto(v) for v in campo("especificacao"))[:200],
                grupo=grupo[:60],
                item=item,
                documento_tipo=" ".join(_texto(v) for v in campo("tipo"))[:30],
                documento_numero="" if normalizar(numero) in _SEM_NUMERO else numero[:40],
                documento_data=next((d for v in campo("data_documento") if (d := _data(v))), None),
                transferencia=next((_digitos(v) for v in campo("transferencia") if _digitos(v)), ""),
                data=next((d for v in campo("data_pagamento") if (d := _data(v))), None),
            )
        )
        soma += valor
    if lido.total_despesas is not None and lido.total_despesas != soma:
        lido.avisos.append(
            Aviso(
                f"A soma dos pagamentos lidos (R$ {_reais(soma)}) não bate com o “TOTAL DAS DESPESAS” "
                f"(R$ {_reais(lido.total_despesas)}): confira se alguma linha ficou fora do lugar."
            )
        )


def _reais(valor: Decimal) -> str:
    return f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
