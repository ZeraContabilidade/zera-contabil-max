"""Perguntas ao cliente: uma planilha com os movimentos que o sistema não reconheceu (“o que é este Pix de R$ 77?”)
para mandar ao cliente. A planilha respondida volta pela tela: a resposta fica na linha da Revisão, entra no histórico
do lançamento (no lugar da natureza) e a IA usa a resposta para sugerir a conta. A escolha final ensina o sistema."""

from dataclasses import dataclass
from datetime import UTC, date, datetime
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.modelos import Empresa, MovimentoExtrato, Sugestao
from app.servicos.competencia import Periodo, como_periodo, nome_competencia
from app.servicos.plano import ErroLeitura

CABECALHO = (
    "Ref.",
    "Data",
    "Conta no banco",
    "Descrição no banco",
    "Valor (R$)",
    "Entrou/Saiu",
    "O que é? (responda aqui)",
)
TAMANHO_RESPOSTA = 200


def linhas_para_perguntar(sessao: Session, empresa: Empresa, competencia: date | Periodo) -> list[Sugestao]:
    """As vermelhas do período: pendentes, sem conferência, que nem as regras nem a IA reconheceram."""
    inicio, fim = como_periodo(competencia).limites
    return list(
        sessao.scalars(
            select(Sugestao)
            .join(MovimentoExtrato, Sugestao.movimento_id == MovimentoExtrato.id)
            .where(
                Sugestao.empresa_id == empresa.id,
                MovimentoExtrato.data.between(inicio, fim),
                Sugestao.status == "PENDENTE",
                Sugestao.confianca < 70,
                Sugestao.conferido_por_id.is_(None),
            )
            .order_by(MovimentoExtrato.data, MovimentoExtrato.conta_bancaria_id, MovimentoExtrato.ordem)
        )
    )


def planilha_perguntas(sessao: Session, empresa: Empresa, competencia: date | Periodo) -> bytes:
    periodo = como_periodo(competencia)
    livro = Workbook()
    folha = livro.active
    folha.title = "Perguntas"
    folha.append([f"{empresa.nome} — {nome_competencia(periodo)}: o que é cada movimento?"])
    folha["A1"].font = Font(bold=True, size=13)
    folha.append(["Escreva na coluna amarela o que foi cada pagamento ou recebimento (ex.: material de limpeza, "
                  "aluguel, conserto da geladeira). Não mude as outras colunas."])  # fmt: skip
    folha["A2"].font = Font(italic=True, color="6B665C")
    folha.append([])
    folha.append(list(CABECALHO))
    amarelo = PatternFill("solid", fgColor="FFF2CC")
    for celula in folha[4]:
        celula.font = Font(bold=True)
    for sugestao in linhas_para_perguntar(sessao, empresa, periodo):
        m = sugestao.movimento
        descricao = " · ".join(p for p in (m.historico, m.complemento) if p)
        valor = -m.valor if m.sentido == "SAIDA" else m.valor
        folha.append([sugestao.id, m.data, m.conta_bancaria.rotulo, descricao, valor,
                      "saiu" if m.sentido == "SAIDA" else "entrou", sugestao.resposta_cliente or ""])  # fmt: skip
        linha = folha.max_row
        folha.cell(linha, 2).number_format = "DD/MM/YYYY"
        folha.cell(linha, 5).number_format = "#,##0.00;[Red]-#,##0.00"
        folha.cell(linha, 7).fill = amarelo
        folha.cell(linha, 4).alignment = Alignment(wrap_text=True, vertical="top")
    for coluna, largura in zip("ABCDEFG", (8, 12, 34, 52, 14, 11, 48), strict=True):
        folha.column_dimensions[coluna].width = largura
    folha.freeze_panes = "A5"
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()


@dataclass
class ResultadoRespostas:
    respondidas: list[Sugestao]
    ja_decididas: int = 0  # a linha já foi conferida ou exportada: a resposta não muda nada
    nao_achadas: int = 0

    def resumo(self) -> str:
        texto = f"Respostas do cliente: {len(self.respondidas)} linha(s) receberam a resposta"
        if self.ja_decididas:
            texto += f", {self.ja_decididas} já tinham sido decididas"
        if self.nao_achadas:
            texto += f", {self.nao_achadas} não são desta empresa (planilha de outra?)"
        return texto + "."


def ler_respostas(sessao: Session, empresa: Empresa, conteudo: bytes) -> ResultadoRespostas:
    """A planilha de perguntas respondida → a resposta em cada linha (só nas ainda não decididas)."""
    try:
        folha = load_workbook(BytesIO(conteudo), read_only=True, data_only=True).active
        linhas = list(folha.iter_rows(values_only=True))
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as erro:
        raise ErroLeitura(
            "Não consegui abrir a planilha. Envie o arquivo .xlsx de perguntas que o sistema gerou."
        ) from erro
    inicio = next((n for n, linha in enumerate(linhas) if linha and linha[0] == CABECALHO[0]), None)
    if inicio is None:
        raise ErroLeitura(
            "Esta planilha não é a de perguntas ao cliente: baixe a planilha pela Revisão e mande de volta."
        )
    resultado = ResultadoRespostas([])
    agora = datetime.now(UTC)
    for linha in linhas[inicio + 1 :]:
        if not linha or not isinstance(linha[0], int | float) or len(linha) < 7:
            continue
        resposta = " ".join(str(linha[6] or "").replace(";", ",").split())[:TAMANHO_RESPOSTA]
        if not resposta:
            continue
        sugestao = sessao.get(Sugestao, int(linha[0]))
        if sugestao is None or sugestao.empresa_id != empresa.id:
            resultado.nao_achadas += 1
            continue
        if sugestao.status == "EXPORTADO" or sugestao.conferido_por_id is not None:
            resultado.ja_decididas += 1
            continue
        sugestao.resposta_cliente = resposta
        sugestao.resposta_em = agora
        sugestao.revisado_ia = False  # a IA olha de novo, agora com a resposta
        resultado.respondidas.append(sugestao)
    return resultado
