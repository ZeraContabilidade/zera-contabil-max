"""O `lanctos.txt` que o Domínio importa (§3.6, §11.1) — igual, byte a byte, ao que a macro `Gerar()` da planilha
EXPORTAÇÃO DOMINIO.xlsm gravava e o Domínio já aceita:

    05/01/2026;492;8;150,17;;QUATAI TRANSPORTE ... - VALE TRANSPORTE - 841;1;230;;

10 campos separados por “;” (data; débito; crédito; valor; cód. histórico vazio; complemento; inicia lote = 1;
código da empresa; c. custo débito e crédito vazios), sem cabeçalho, quebra de linha CRLF, Windows-1252. Valor com
vírgula, sem milhar e sem zeros sobrando (346,5 · 120). Conta em branco é aceita pelo Domínio (decisão de 08/10).

`conferir_txt` confere um arquivo pronto linha a linha (o modelo do escritório tinha 3 erros: “28/01//2026”,
“28/01//2027” e a conta “10595.”) — o sistema confere o que ele mesmo gera antes de entregar.
"""

import re
import unicodedata
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.auditoria import registrar_evento
from app.modelos import ContaBancaria, ContaContabil, Empresa, Exportacao, MovimentoExtrato, Sugestao, Usuario
from app.servicos.armazenamento import guardar_arquivo
from app.servicos.competencia import Periodo, como_periodo, nome_competencia, periodo_no_arquivo
from app.servicos.extratos import limites_do_mes

SEPARADOR = ";"
QUEBRA = "\r\n"
CODIFICACAO = "cp1252"
CAMPOS = 10
_DATA = re.compile(r"^\d{2}/\d{2}/\d{4}$")
_CONTA = re.compile(r"^\d+$")
_VALOR = re.compile(r"^\d+(,\d{1,2})?$")


@dataclass(frozen=True)
class LinhaLancto:
    data: date
    debito: int | None  # código reduzido; None = em branco (conta fora do plano — completa no Domínio)
    credito: int | None
    valor: Decimal
    complemento: str


def formatar_valor(valor: Decimal) -> str:
    """Como a macro gravava: vírgula decimal, sem milhar e sem zeros sobrando (150,17 · 346,5 · 120)."""
    inteiro, _, centavos = f"{valor.quantize(Decimal('0.01')):f}".partition(".")
    centavos = centavos.rstrip("0")
    return f"{inteiro},{centavos}" if centavos else inteiro


def texto_windows(texto: str) -> str:
    """Só o que o Windows-1252 grava (letra fora dele perde o acento ou some); sem quebra de linha nem tabulação."""
    saida = []
    for letra in " ".join((texto or "").replace("\t", " ").splitlines()):
        try:
            letra.encode(CODIFICACAO)
            saida.append(letra)
        except UnicodeEncodeError:
            saida.append(unicodedata.normalize("NFKD", letra).encode(CODIFICACAO, "ignore").decode(CODIFICACAO))
    return "".join(saida)


def linha_txt(linha: LinhaLancto, codigo_empresa: int) -> str:
    campos = [
        f"{linha.data:%d/%m/%Y}",
        str(linha.debito or ""),
        str(linha.credito or ""),
        formatar_valor(linha.valor),
        "",  # código do histórico (não usado)
        texto_windows(linha.complemento),
        "1",  # inicia lote
        str(codigo_empresa),
        "",  # centro de custo débito
        "",  # centro de custo crédito
    ]
    return SEPARADOR.join(campos)


def gerar_txt(linhas: list[LinhaLancto], codigo_empresa: int) -> bytes:
    return "".join(linha_txt(linha, codigo_empresa) + QUEBRA for linha in linhas).encode(CODIFICACAO)


def ler_txt(conteudo: bytes) -> list[LinhaLancto]:
    """Lê um lanctos.txt já conferido (sem erros) de volta em linhas — usado nos testes do arquivo de ouro."""
    linhas = []
    for texto in conteudo.decode(CODIFICACAO).split(QUEBRA):
        if not texto:
            continue
        campos = texto.split(SEPARADOR)
        linhas.append(
            LinhaLancto(
                datetime.strptime(campos[0], "%d/%m/%Y").date(),
                int(campos[1]) if campos[1] else None,
                int(campos[2]) if campos[2] else None,
                Decimal(campos[3].replace(",", ".")),
                campos[5],
            )
        )
    return linhas


# ---------------------------------------------------------------- conferência do arquivo pronto


@dataclass(frozen=True)
class Problema:
    linha: int  # nº da linha no arquivo (a 1ª é 1)
    texto: str


def _conferir_conta(numero: int, lado: str, campo: str, contas: dict[int, ContaContabil] | None) -> str | None:
    if not campo:
        return None  # em branco: o Domínio aceita
    if not _CONTA.match(campo):
        return f"Linha {numero}: conta {lado} “{campo}” inválida (só o código reduzido, ex.: 355)."
    if contas is None:
        return None
    conta = contas.get(int(campo))
    if conta is None or not conta.ativa:
        return f"Linha {numero}: a conta {lado} {campo} não existe no plano de contas desta empresa."
    if not conta.analitica:
        return f"Linha {numero}: a conta {lado} {conta.rotulo} é sintética (grupo): use uma conta analítica."
    return None


def conferir_txt(
    conteudo: bytes,
    contas: dict[int, ContaContabil] | None = None,
    inicio: date | None = None,
    fim: date | None = None,
    codigo_empresa: int | None = None,
) -> list[Problema]:
    """Confere um lanctos.txt linha a linha. `contas` = plano da empresa por código reduzido (None = não confere)."""
    try:
        texto = conteudo.decode(CODIFICACAO)
    except UnicodeDecodeError:
        return [Problema(0, "O arquivo não está em Windows-1252 (o Domínio não lê os acentos).")]
    problemas = []
    if "\n" in texto.replace(QUEBRA, "") or (texto and not texto.endswith(QUEBRA)):
        problemas.append(Problema(0, "O arquivo tem de usar quebra de linha do Windows (CRLF), inclusive na última."))
    for numero, linha in enumerate(texto.replace(QUEBRA, "\n").split("\n"), 1):
        if not linha:
            continue
        campos = linha.split(SEPARADOR)
        if len(campos) != CAMPOS:
            problemas.append(
                Problema(
                    numero, f"Linha {numero}: tem {len(campos)} campos em vez de {CAMPOS} (há “;” no complemento?)."
                )
            )
            continue
        data, debito, credito, valor, _, _, lote, codigo, _, _ = campos
        erros = []
        quando = None
        if _DATA.match(data):
            try:
                quando = datetime.strptime(data, "%d/%m/%Y").date()
            except ValueError:
                pass
        if quando is None:
            erros.append(f"Linha {numero}: data “{data}” inválida (use dd/mm/aaaa).")
        elif (inicio and quando < inicio) or (fim and quando > fim):
            erros.append(f"Linha {numero}: data {data} fora do período exportado.")
        erros += [e for e in (_conferir_conta(numero, "débito", debito, contas),
                              _conferir_conta(numero, "crédito", credito, contas)) if e]  # fmt: skip
        if not debito and not credito:
            erros.append(f"Linha {numero}: débito e crédito em branco.")
        elif debito and debito == credito:
            erros.append(f"Linha {numero}: débito e crédito na mesma conta ({debito}).")
        try:
            valido = bool(_VALOR.match(valor)) and Decimal(valor.replace(",", ".")) > 0
        except InvalidOperation:
            valido = False
        if not valido:
            erros.append(f"Linha {numero}: valor “{valor}” inválido (maior que zero, vírgula decimal, ex.: 150,17).")
        if lote != "1":
            erros.append(f"Linha {numero}: o campo “inicia lote” tem de ser 1.")
        if codigo_empresa is not None and codigo != str(codigo_empresa):
            erros.append(f"Linha {numero}: código da empresa {codigo} diferente de {codigo_empresa}.")
        problemas += [Problema(numero, erro) for erro in erros]
    return problemas


# ---------------------------------------------------------------- exportar o mês (§3.6)


class ErroExportacao(Exception):
    """Mensagem em português dizendo o que fazer."""


@dataclass(frozen=True)
class ProblemaLinha:
    sugestao: Sugestao
    texto: str


@dataclass
class TotalConta:
    rotulo: str  # “355 – ÁGUA E ESGOTO” ou “(em branco)”
    debitos: Decimal = Decimal("0")
    creditos: Decimal = Decimal("0")
    linhas: int = 0


@dataclass
class Preparo:
    periodo: Periodo  # um mês ou vários (de um mês até outro)
    disponiveis: list[ContaBancaria]  # contas com movimento no período
    escolhidas: list[ContaBancaria]
    prontas: list[Sugestao] = field(default_factory=list)  # aprovadas: vão para o arquivo
    pendentes: list[Sugestao] = field(default_factory=list)
    exportadas: int = 0
    sem_lancamento: int = 0
    problemas: list[ProblemaLinha] = field(default_factory=list)

    @property
    def total(self) -> Decimal:
        return sum((s.movimento.valor for s in self.prontas), Decimal("0"))

    @property
    def por_conta(self) -> list[TotalConta]:
        return totais_por_conta(self.prontas)


def totais_por_conta(sugestoes: list[Sugestao]) -> list[TotalConta]:
    totais: dict[tuple[int, str], TotalConta] = {}
    for sugestao in sugestoes:
        for conta, lado in ((sugestao.conta_debito, "debitos"), (sugestao.conta_credito, "creditos")):
            chave = (conta.reduzido, conta.rotulo) if conta else (10**9, "(em branco — completar no Domínio)")
            total = totais.setdefault(chave, TotalConta(chave[1]))
            setattr(total, lado, getattr(total, lado) + sugestao.movimento.valor)
            total.linhas += 1
    return [totais[chave] for chave in sorted(totais)]


def contas_do_mes(sessao: Session, empresa: Empresa, competencia: date | Periodo) -> list[ContaBancaria]:
    inicio, fim = limites_do_mes(competencia)
    com_movimento = select(MovimentoExtrato.conta_bancaria_id).where(
        MovimentoExtrato.empresa_id == empresa.id, MovimentoExtrato.data.between(inicio, fim)
    )
    return list(
        sessao.scalars(
            select(ContaBancaria)
            .where(ContaBancaria.id.in_(com_movimento))
            .order_by(ContaBancaria.banco_nome, ContaBancaria.numero, ContaBancaria.tipo)
        )
    )


def _problemas(sugestao: Sugestao) -> list[str]:
    erros = []
    for lado, conta in (("débito", sugestao.conta_debito), ("crédito", sugestao.conta_credito)):
        if conta is not None and not conta.ativa:
            erros.append(f"a conta {lado} {conta.rotulo} saiu do plano de contas: escolha outra na Revisão.")
        elif conta is not None and not conta.analitica:
            erros.append(f"a conta {lado} {conta.rotulo} é sintética (grupo): escolha uma analítica na Revisão.")
    if sugestao.conta_debito is None and sugestao.conta_credito is None:
        erros.append("sem débito e sem crédito: escolha a conta na Revisão.")
    elif sugestao.conta_debito is not None and sugestao.conta_debito_id == sugestao.conta_credito_id:
        erros.append(f"débito e crédito na mesma conta ({sugestao.conta_debito.rotulo}): corrija na Revisão.")
    if sugestao.movimento.valor <= 0:
        erros.append("valor zero ou negativo.")
    if SEPARADOR in (sugestao.complemento or ""):
        erros.append("o complemento tem “;” (o Domínio usa “;” para separar os campos): troque na Revisão.")
    return erros


def preparar(
    sessao: Session, empresa: Empresa, competencia: date | Periodo, contas_ids: list[int] | None = None
) -> Preparo:
    """O que vai para o arquivo (aprovadas), o que fica de fora (pendentes) e o que impede a exportação —
    de um mês ou de um período inteiro (ex.: o ano, para importar no Domínio de uma vez)."""
    periodo = como_periodo(competencia)
    inicio, fim = periodo.limites
    disponiveis = contas_do_mes(sessao, empresa, periodo)
    escolhidas = [c for c in disponiveis if not contas_ids or c.id in contas_ids]
    preparo = Preparo(periodo, disponiveis, escolhidas)
    if not escolhidas:
        return preparo
    sugestoes = sessao.scalars(
        select(Sugestao)
        .join(MovimentoExtrato, Sugestao.movimento_id == MovimentoExtrato.id)
        .where(
            Sugestao.empresa_id == empresa.id,
            MovimentoExtrato.data.between(inicio, fim),
            MovimentoExtrato.conta_bancaria_id.in_([c.id for c in escolhidas]),
        )
        .order_by(MovimentoExtrato.data, MovimentoExtrato.conta_bancaria_id, MovimentoExtrato.ordem)
    )
    for sugestao in sugestoes:
        if sugestao.status == "APROVADO":
            preparo.prontas.append(sugestao)
            preparo.problemas += [ProblemaLinha(sugestao, texto) for texto in _problemas(sugestao)]
        elif sugestao.status == "PENDENTE":
            preparo.pendentes.append(sugestao)
        elif sugestao.status == "EXPORTADO":
            preparo.exportadas += 1
        else:
            preparo.sem_lancamento += 1
    return preparo


def linha_da_sugestao(sugestao: Sugestao) -> LinhaLancto:
    movimento = sugestao.movimento
    return LinhaLancto(
        movimento.data,
        sugestao.conta_debito.reduzido if sugestao.conta_debito else None,
        sugestao.conta_credito.reduzido if sugestao.conta_credito else None,
        movimento.valor,
        sugestao.complemento or movimento.descricao,
    )


def exportar(
    sessao: Session,
    empresa: Empresa,
    competencia: date | Periodo,
    contas_ids: list[int] | None,
    usuario: Usuario | None,
    deixar_pendentes_de_fora: bool = False,
    automatica: bool = False,
) -> Exportacao:
    """Gera o lanctos.txt do mês (ou do período), guarda o arquivo, registra e trava as linhas exportadas."""
    periodo = como_periodo(competencia)
    preparo = preparar(sessao, empresa, periodo, contas_ids)
    mes = nome_competencia(periodo)
    if not preparo.prontas:
        raise ErroExportacao(f"Nada aprovado para exportar em {mes}: confira a aba Revisão.")
    if preparo.problemas:
        raise ErroExportacao(
            f"{len(preparo.problemas)} problema(s) impedem a exportação: corrija as linhas mostradas e tente de novo."
        )
    if preparo.pendentes and not deixar_pendentes_de_fora:
        raise ErroExportacao(
            f"Há {len(preparo.pendentes)} linha(s) pendente(s) em {mes}: aprove na Revisão ou marque "
            "“Deixar as pendentes de fora”."
        )
    inicio, fim = periodo.limites
    conteudo = gerar_txt([linha_da_sugestao(s) for s in preparo.prontas], empresa.codigo_no_lanctos)
    plano = {c.reduzido: c for c in sessao.scalars(select(ContaContabil).where(ContaContabil.empresa_id == empresa.id))}
    if falhas := conferir_txt(conteudo, plano, inicio, fim, empresa.codigo_no_lanctos):
        raise ErroExportacao(f"O arquivo gerado não passou na conferência ({falhas[0].texto}). Avise o suporte.")
    nome = f"lanctos_{empresa.codigo_arquivo}_{periodo_no_arquivo(periodo)}.txt"
    arquivo = guardar_arquivo(sessao, empresa, "EXPORTACAO", nome, conteudo, usuario)
    exportacao = Exportacao(
        empresa_id=empresa.id,
        competencia=periodo.inicio,
        competencia_fim=None if periodo.um_mes else periodo.fim,
        contas_bancarias=[c.id for c in preparo.escolhidas],
        linhas=len(preparo.prontas),
        total=preparo.total,
        deixadas_de_fora=len(preparo.pendentes),
        arquivo_id=arquivo.id,
        usuario_id=usuario.id if usuario else None,
        automatica=automatica,
    )
    sessao.add(exportacao)
    sessao.flush()
    sessao.info["auditoria_em_bloco"] = True  # um evento no log, não uma linha por sugestão
    try:
        for sugestao in preparo.prontas:
            sugestao.status = "EXPORTADO"
            sugestao.exportacao_id = exportacao.id
        sessao.flush()
    finally:
        sessao.info.pop("auditoria_em_bloco", None)
    fora = f", {len(preparo.pendentes)} pendente(s) de fora" if preparo.pendentes else ""
    total = formatar_valor(exportacao.total)
    detalhe = f"{empresa.nome}, {mes}: {exportacao.linhas} lançamento(s), R$ {total}{fora} — {nome}"
    if automatica:
        detalhe += " (exportação automática)"
    registrar_evento(sessao, "EXPORTAR", usuario.id if usuario else None, detalhe)
    return exportacao


def reabrir_exportacao(sessao: Session, exportacao: Exportacao, usuario: Usuario, motivo: str) -> int:
    """Só o Supervisor (a rota exige): as linhas daquela exportação voltam a APROVADO, para corrigir e exportar de
    novo. O arquivo antigo continua guardado, marcado como reaberto (não importar). Devolve quantas linhas voltaram."""
    motivo = " ".join((motivo or "").split())
    if exportacao.reaberta_em is not None:
        raise ErroExportacao("Esta exportação já foi reaberta.")
    if len(motivo) < 5:
        raise ErroExportacao("Escreva o motivo da reabertura (ele fica registrado no log).")
    sugestoes = list(
        sessao.scalars(select(Sugestao).where(Sugestao.exportacao_id == exportacao.id, Sugestao.status == "EXPORTADO"))
    )
    sessao.info["auditoria_em_bloco"] = True  # um evento no log, não uma linha por sugestão
    try:
        for sugestao in sugestoes:
            sugestao.status = "APROVADO"
            sugestao.exportacao_id = None
        exportacao.reaberta_em = datetime.now(UTC)
        exportacao.reaberta_por_id = usuario.id
        exportacao.motivo_reabertura = motivo[:2000]
        sessao.flush()
    finally:
        sessao.info.pop("auditoria_em_bloco", None)
    empresa = sessao.get(Empresa, exportacao.empresa_id)
    detalhe = (
        f"{empresa.nome}, {nome_competencia(exportacao.periodo)}: {len(sugestoes)} linha(s) reaberta(s) "
        f"({exportacao.arquivo.nome_original}). Motivo: {motivo}"
    )
    registrar_evento(sessao, "REABRIR_EXPORTACAO", usuario.id, detalhe)
    return len(sugestoes)


def exportacoes(sessao: Session, empresa: Empresa) -> list[Exportacao]:
    return list(
        sessao.scalars(
            select(Exportacao).where(Exportacao.empresa_id == empresa.id).order_by(Exportacao.criado_em.desc())
        )
    )


def situacao_do_mes(sessao: Session, empresa: Empresa, competencia: date | Periodo) -> tuple[int, int]:
    """(lançamentos exportados no mês ou período, aprovados que ainda não foram) — para a barra de progresso."""
    inicio, fim = limites_do_mes(competencia)
    contagem = dict(
        sessao.execute(
            select(Sugestao.status, func.count())
            .join(MovimentoExtrato, Sugestao.movimento_id == MovimentoExtrato.id)
            .where(Sugestao.empresa_id == empresa.id, MovimentoExtrato.data.between(inicio, fim))
            .group_by(Sugestao.status)
        ).all()
    )
    return contagem.get("EXPORTADO", 0), contagem.get("APROVADO", 0)


def resumo_excel(sessao: Session, empresa: Empresa, exportacao: Exportacao) -> bytes:
    """Resumo da exportação: total por conta + todas as linhas (para conferir com o Domínio depois de importar)."""
    sugestoes = list(
        sessao.scalars(
            select(Sugestao)
            .join(MovimentoExtrato, Sugestao.movimento_id == MovimentoExtrato.id)
            .where(Sugestao.exportacao_id == exportacao.id)
            .order_by(MovimentoExtrato.data, MovimentoExtrato.conta_bancaria_id, MovimentoExtrato.ordem)
        )
    )
    livro = Workbook()
    folha = livro.active
    folha.title = "Resumo por conta"
    folha.append([f"{empresa.nome} (código {empresa.rotulo_codigo}) — {nome_competencia(exportacao.periodo)}"])
    folha["A1"].font = Font(bold=True, size=13)
    folha.append([f"{exportacao.linhas} lançamento(s) exportado(s) em {exportacao.criado_em:%d/%m/%Y %H:%M}"])
    folha.append([])
    folha.append(["Conta", "Débitos (R$)", "Créditos (R$)", "Linhas"])
    for celula in folha[4]:
        celula.font = Font(bold=True)
    for total in totais_por_conta(sugestoes):
        folha.append([total.rotulo, total.debitos, total.creditos, total.linhas])
    for linha in folha.iter_rows(min_row=5):
        linha[1].number_format = linha[2].number_format = "#,##0.00"
    for coluna, largura in zip("ABCD", (52, 16, 16, 9), strict=True):
        folha.column_dimensions[coluna].width = largura

    detalhe = livro.create_sheet("Lançamentos")
    detalhe.append(["Data", "Débito", "Crédito", "Valor (R$)", "Complemento", "Conta bancária"])
    for celula in detalhe[1]:
        celula.font = Font(bold=True)
    for sugestao in sugestoes:
        linha = linha_da_sugestao(sugestao)
        detalhe.append(
            [linha.data, linha.debito, linha.credito, linha.valor, linha.complemento,
             sugestao.movimento.conta_bancaria.rotulo]
        )  # fmt: skip
    for linha in detalhe.iter_rows(min_row=2):
        linha[0].number_format = "DD/MM/YYYY"
        linha[3].number_format = "#,##0.00"
    for coluna, largura in zip("ABCDEF", (12, 9, 9, 14, 70, 48), strict=True):
        detalhe.column_dimensions[coluna].width = largura
    saida = BytesIO()
    livro.save(saida)
    return saida.getvalue()
