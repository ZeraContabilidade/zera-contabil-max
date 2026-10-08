"""Aba Exportar: o que vai e o que fica de fora, pendentes só com confirmação, problemas que bloqueiam mostrando a
linha, o lanctos.txt no formato do Domínio (conferido), resumo em Excel, registro e linhas travadas (§3.6)."""

from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.modelos import (
    ArquivoOriginal,
    ContaBancaria,
    ContaContabil,
    Exportacao,
    Log,
    MovimentoExtrato,
    Papel,
    Sugestao,
)
from app.servicos.classificacao import classificar
from app.servicos.exportacao import conferir_txt, preparar
from app.servicos.extratos import importar_extrato
from app.servicos.progresso import progresso_do_mes
from tests.conftest import entrar
from tests.test_classificacao import POUPANCA, extrato
from tests.test_revisao import (  # noqa: F401 — mesma empresa de teste (histórico, contas-chave e extrato de agosto)
    AGOSTO,
    empresa,
    recarregar,
    sugestao_de,
)

pytestmark = pytest.mark.banco
MES = {"mes": "2026-08"}


def gerar(cliente, empresa, **dados):  # noqa: F811
    return cliente.post(f"/empresas/{empresa.id}/exportar", data={**MES, **dados}, follow_redirects=True)


def aprovadas(sessao) -> list[Sugestao]:
    return list(sessao.scalars(select(Sugestao).where(Sugestao.status == "APROVADO")))


def test_tela_mostra_o_que_vai_e_o_que_fica_de_fora(cliente, empresa, sessao):  # noqa: F811
    tela = cliente.get(f"/empresas/{empresa.id}", params={"aba": "exportar", **MES}).text
    prontas = len(aprovadas(sessao))
    assert f"✓ {prontas} lançamento(s) prontos" in tela and "pendente(s) ficam de fora" in tela
    assert "Banco do Brasil · ag. 1234-5 · conta corrente 12345-6" in tela and "Gerar lanctos.txt" in tela
    assert "Deixar as 5 linha(s) pendente(s) de fora deste arquivo" in tela
    assert "Total por conta" in tela and "355 – ÁGUA E ESGOTO" in tela


def test_pendentes_so_ficam_de_fora_com_confirmacao(cliente, empresa, sessao):  # noqa: F811
    tela = gerar(cliente, empresa).text
    assert "Há 5 linha(s) pendente(s) em agosto/2026: aprove na Revisão ou marque" in tela
    assert sessao.scalar(select(Exportacao)) is None


def test_gera_o_arquivo_conferido_e_trava_as_linhas(cliente, empresa, sessao):  # noqa: F811
    salario = sugestao_de(sessao, "MARIA SOUZA")
    salario.conta_credito = None  # conta do banco fora do plano: vai em branco (o Domínio aceita)
    sessao.commit()
    prontas = len(aprovadas(sessao))

    tela = gerar(cliente, empresa, deixar_pendentes="on").text
    assert f"lanctos.txt gerado: {prontas} lançamento(s)." in tela and "Arquivos gerados" in tela
    exportacao = sessao.scalar(select(Exportacao))
    assert (exportacao.linhas, exportacao.deixadas_de_fora, exportacao.competencia) == (prontas, 5, AGOSTO)
    assert sessao.get(ArquivoOriginal, exportacao.arquivo_id).tipo == "EXPORTACAO"

    resposta = cliente.get(f"/empresas/{empresa.id}/exportacoes/{exportacao.id}/lanctos.txt")
    assert 'filename="lanctos_230_2026-08.txt"' in resposta.headers["content-disposition"]
    conteudo = resposta.content
    assert conferir_txt(conteudo, inicio=AGOSTO, fim=date(2026, 8, 31), codigo_empresa=230) == []
    linhas = conteudo.decode("cp1252").split("\r\n")
    assert len(linhas) == prontas + 1  # + o vazio depois do último CRLF
    assert linhas[0].startswith("03/08/2026;355;8;98,1;;COMPANHIA DE AGUA EXEMPLO - 44.555.666/0001-81 - ")
    assert linhas[0].endswith(";1;230;;")
    assert any(linha.startswith("06/08/2026;187;;2861,25;;") for linha in linhas)  # crédito em branco

    salario = recarregar(sessao, salario)
    assert (salario.status, salario.exportacao_id) == ("EXPORTADO", exportacao.id)
    classificar(sessao, empresa, AGOSTO, date(2026, 8, 31))  # a classificação não mexe no exportado
    assert recarregar(sessao, salario).status == "EXPORTADO"
    tela = cliente.post(f"/empresas/{empresa.id}/revisao/{salario.id}", data={"acao": "ignorar"}, follow_redirects=True)
    assert "Esta linha não pode mais ser mudada" in tela.text
    assert sessao.scalar(select(Log).where(Log.acao == "EXPORTAR")).detalhe.startswith("Creche Teste, agosto/2026:")
    etapa = next(e for e in progresso_do_mes(sessao, empresa, AGOSTO) if e.chave == "exportado")
    assert etapa.situacao == "feita"

    assert "Nada aprovado para exportar em agosto/2026" in gerar(cliente, empresa, deixar_pendentes="on").text
    fulano = sugestao_de(sessao, "FULANO DESCONHECIDO")  # resolvida depois: vai num 2º arquivo
    cliente.post(f"/empresas/{empresa.id}/revisao/{fulano.id}", data={"acao": "salvar", "conta": "359"})
    gerar(cliente, empresa, deixar_pendentes="on")
    segunda = sessao.scalars(select(Exportacao).order_by(Exportacao.id.desc())).first()
    assert (segunda.linhas, segunda.deixadas_de_fora) == (1, 4)


def test_problemas_bloqueiam_e_mostram_a_linha(cliente, empresa, sessao):  # noqa: F811
    agua = sugestao_de(sessao, "COMPANHIA DE AGUA EXEMPL")
    agua.complemento = "COMPANHIA DE AGUA; CONTA DE AGUA"
    salario = sugestao_de(sessao, "MARIA SOUZA")
    sessao.scalar(select(ContaContabil).where(ContaContabil.reduzido == 187)).ativa = False  # saiu do plano
    sessao.commit()
    tela = cliente.get(f"/empresas/{empresa.id}", params={"aba": "exportar", **MES}).text
    assert "2 problema(s) impedem a exportação." in tela and "Gerar lanctos.txt" not in tela
    assert "o complemento tem “;”" in tela and "a conta débito 187 – SALÁRIOS E ORDENADOS A PAGAR saiu do plano" in tela
    assert "2 problema(s) impedem a exportação: corrija" in gerar(cliente, empresa, deixar_pendentes="on").text
    assert sessao.scalar(select(Exportacao)) is None and recarregar(sessao, salario).status == "APROVADO"


def test_escolher_as_contas_bancarias(cliente, empresa, sessao, monkeypatch):  # noqa: F811
    monkeypatch.setattr("app.servicos.extratos.ler_extrato", lambda nome, conteudo: extrato("POUPANCA", POUPANCA))
    importar_extrato(sessao, empresa, "poupanca.pdf", b"%PDF-falso", None)
    sessao.commit()
    corrente, poupanca = sessao.scalars(select(ContaBancaria).order_by(ContaBancaria.tipo)).all()
    assert (corrente.tipo, poupanca.tipo) == ("BANCO", "POUPANCA")
    so_corrente = preparar(sessao, empresa, AGOSTO, [corrente.id])
    todas = preparar(sessao, empresa, AGOSTO, [])
    assert len(todas.disponiveis) == 2 and len(so_corrente.escolhidas) == 1
    assert {s.movimento.conta_bancaria_id for s in so_corrente.prontas} == {corrente.id}

    def linhas(preparo):
        return len(preparo.prontas) + len(preparo.pendentes) + preparo.sem_lancamento

    assert linhas(todas) == linhas(so_corrente) + len(POUPANCA)
    tela = cliente.get(f"/empresas/{empresa.id}", params={"aba": "exportar", "bancos": corrente.id, **MES}).text
    assert f'value="{corrente.id}" checked' in tela and f'value="{poupanca.id}" >' in tela


def test_resumo_em_excel(cliente, empresa, sessao):  # noqa: F811
    gerar(cliente, empresa, deixar_pendentes="on")
    exportacao = sessao.scalar(select(Exportacao))
    resposta = cliente.get(f"/empresas/{empresa.id}/exportacoes/{exportacao.id}/resumo.xlsx")
    assert "resumo_230_2026-08.xlsx" in resposta.headers["content-disposition"]
    livro = load_workbook(BytesIO(resposta.content))
    resumo = {linha[0]: linha[1:] for linha in livro["Resumo por conta"].iter_rows(min_row=5, values_only=True)}
    assert float(resumo["355 – ÁGUA E ESGOTO"][0]) == 98.10  # débitos
    linhas = list(livro["Lançamentos"].iter_rows(min_row=2, values_only=True))
    assert len(linhas) == exportacao.linhas
    movimento = sessao.scalar(
        select(MovimentoExtrato).where(MovimentoExtrato.complemento == "COMPANHIA DE AGUA EXEMPL")
    )
    assert linhas[0][:3] == (movimento.data, 355, 8) or linhas[0][0].date() == movimento.data


def test_so_o_supervisor_reabre_a_exportacao(cliente, empresa, sessao, criar_usuario):  # noqa: F811
    gerar(cliente, empresa, deixar_pendentes="on")
    exportacao = sessao.scalar(select(Exportacao))
    linhas = exportacao.linhas
    reabrir = f"/empresas/{empresa.id}/exportacoes/{exportacao.id}/reabrir"

    criar_usuario("carla", papel=Papel.ASSISTENTE)
    entrar(cliente, "carla")
    tela = cliente.get(f"/empresas/{empresa.id}", params={"aba": "exportar", **MES}).text
    assert "Arquivos gerados" in tela and "Reabrir esta exportação" not in tela  # o assistente nem vê o botão
    assert cliente.post(reabrir, data={"motivo": "quero mudar"}).status_code == 403
    sessao.refresh(exportacao)
    assert exportacao.reaberta_em is None

    entrar(cliente, "ana")  # supervisora
    tela = cliente.get(f"/empresas/{empresa.id}", params={"aba": "exportar", **MES}).text
    assert "Reabrir esta exportação" in tela and "exclua lá os lançamentos dele" in tela
    assert "Escreva o motivo da reabertura" in cliente.post(reabrir, data={"motivo": " "}, follow_redirects=True).text
    tela = cliente.post(reabrir, data={"motivo": "Conta errada no salário"}, follow_redirects=True).text
    assert f"Exportação reaberta: {linhas} linha(s) voltaram para a Revisão" in tela
    assert "reaberta — não importar" in tela and "Conta errada no salário" in tela

    sessao.expire_all()
    exportacao = sessao.get(Exportacao, exportacao.id)
    assert exportacao.reaberta_por.login == "ana" and exportacao.motivo_reabertura == "Conta errada no salário"
    salario = sugestao_de(sessao, "MARIA SOUZA")
    assert (salario.status, salario.exportacao_id) == ("APROVADO", None)  # pode ser corrigida de novo
    log = sessao.scalar(select(Log).where(Log.acao == "REABRIR_EXPORTACAO"))
    assert log.detalhe.endswith("Motivo: Conta errada no salário") and f"{linhas} linha(s) reaberta(s)" in log.detalhe
    etapa = next(e for e in progresso_do_mes(sessao, empresa, AGOSTO) if e.chave == "exportado")
    assert etapa.situacao != "feita"  # há aprovadas fora de um arquivo válido
    assert "já foi reaberta" in cliente.post(reabrir, data={"motivo": "de novo"}, follow_redirects=True).text

    cliente.post(f"/empresas/{empresa.id}/revisao/{salario.id}", data={"acao": "salvar", "conta": "331"})
    gerar(cliente, empresa, deixar_pendentes="on")  # o arquivo novo sai com a correção
    nova = sessao.scalars(select(Exportacao).order_by(Exportacao.id.desc())).first()
    assert nova.id != exportacao.id and nova.linhas == linhas
    conteudo = cliente.get(f"/empresas/{empresa.id}/exportacoes/{nova.id}/lanctos.txt").content.decode("cp1252")
    assert "06/08/2026;331;8;2861,25;;" in conteudo
